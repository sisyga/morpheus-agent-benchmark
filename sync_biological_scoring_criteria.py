"""Synchronize biological_scoring_criteria.json with the biological scoring criteria workbook.

The workbook has one sheet per paper, named by its canonical paper stem:

    A1   "<stem>: biological scoring criteria v<N>"
    A3   Target                           B3  target text
    A4   Scoring instructions             B4  instructions for the scoring agent
    A5   Workbook formula (qualitative)   B5  "1/n*( clause + clause + ... )"
    A6   Workbook formula (quantitative)  B6  same for numeric clauses
    A8   "#" header, then one row per clause: #, Clause ID, Kind, Weight, Share of paper score, Criterion

The clause table is authoritative, including each clause's weight; the scorer awards a paper the passed weight
divided by the total weight. The Share column must equal weight / total weight and the formula rows must restate
the table's criteria. Any disagreement is reported and nothing is written, so the workbook and the registry cannot
silently drift apart.
"""

import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook


DEFAULT_WORKBOOK = Path("Biological_scoring_criteria_v7.xlsx")
DEFAULT_REGISTRY = Path("biological_scoring_criteria.json")
TITLE_PATTERN = re.compile(r"^(?P<stem>\S+): biological scoring criteria v(?P<version>\d+)$")
CLAUSE_ID_PATTERN = re.compile(r"^[a-z0-9]+(?:_[a-z0-9]+)*$")
KINDS = ("qualitative", "numeric")
FORMULA_LABELS = {
    "qualitative": "Workbook formula (qualitative)",
    "numeric": "Workbook formula (quantitative)",
}
SHARE_TOLERANCE = 1e-4
# Numbers adds this sheet to every xlsx export; it holds no criteria.
IGNORED_SHEETS = {"Export Summary"}
FORMULA_PATTERN = re.compile(r"^1\s*/\s*(?P<count>\d+)\s*\*\s*\(\s*(?P<body>.*?)\s*\)$", re.DOTALL)


class CriteriaWorkbookError(ValueError):
    def __init__(self, problems: list[str]):
        super().__init__("Criteria workbook is inconsistent:\n" + "\n".join(f"- {problem}" for problem in problems))
        self.problems = problems


def criterion_formula(descriptions: list[str]) -> str | None:
    if not descriptions:
        return None
    return f"1/{len(descriptions)}*( " + " + ".join(descriptions) + " )"


def read_criteria_workbook(workbook_path: str | Path) -> tuple[int, list[dict[str, object]]]:
    """Return the workbook version and one criterion per sheet, or raise CriteriaWorkbookError."""
    workbook = load_workbook(workbook_path, data_only=True)
    problems: list[str] = []
    versions: set[int] = set()
    criteria: list[dict[str, object]] = []
    for sheet in workbook.worksheets:
        if sheet.title in IGNORED_SHEETS:
            continue
        criterion = _read_sheet(sheet, problems)
        if criterion is not None:
            versions.add(criterion.pop("version"))
            criteria.append(criterion)
    if len(versions) > 1:
        problems.append(f"sheets disagree on the criteria version: {sorted(versions)}")
    if not criteria and not problems:
        problems.append("workbook contains no criteria sheets")
    if problems:
        raise CriteriaWorkbookError(problems)
    return versions.pop(), criteria


def _read_sheet(sheet, problems: list[str]) -> dict[str, object] | None:
    name = sheet.title
    title = _text(sheet["A1"].value)
    match = TITLE_PATTERN.match(title)
    if match is None:
        problems.append(f"{name}: A1 must read '<stem>: biological scoring criteria v<N>', found {title!r}")
        return None
    if match.group("stem") != name:
        problems.append(f"{name}: A1 names {match.group('stem')!r}, which differs from the sheet name")

    labelled = {_text(row[0].value): (row[1].value if len(row) > 1 else None) for row in sheet.iter_rows(max_col=2)}
    target = _text(labelled.get("Target"))
    instructions = _text(labelled.get("Scoring instructions"))
    for label, value in (("Target", target), ("Scoring instructions", instructions)):
        if not value:
            problems.append(f"{name}: '{label}' is empty")

    clauses = _read_clause_table(sheet, problems)
    if clauses is None:
        return None
    listed_shares = [clause.pop("listed_share") for clause in clauses]
    if all(_is_positive_number(clause["weight"]) for clause in clauses):
        total_weight = sum(clause["weight"] for clause in clauses)
        for clause, listed_share in zip(clauses, listed_shares):
            share = clause["weight"] / total_weight
            if not _is_number(listed_share) or abs(listed_share - share) > SHARE_TOLERANCE:
                problems.append(f"{name}: clause {clause['id']} share is {listed_share!r}, but its weight gives {share:.6g}")

    for kind, label in FORMULA_LABELS.items():
        expected = [_normalized(clause["description"]) for clause in clauses if clause["kind"] == kind]
        if formula_clauses(_text(labelled.get(label))) != expected:
            problems.append(f"{name}: '{label}' does not restate the {kind} clauses of the clause table")

    return {
        "version": int(match.group("version")),
        "paper_stem": name,
        "target": target,
        "scoring_instructions": instructions,
        "clauses": clauses,
    }


def _read_clause_table(sheet, problems: list[str]) -> list[dict[str, object]] | None:
    name = sheet.title
    rows = list(sheet.iter_rows(values_only=True))
    header_index = next((index for index, row in enumerate(rows) if _text(row[0] if row else None) == "#"), None)
    if header_index is None:
        problems.append(f"{name}: no clause table header row starting with '#'")
        return None
    clauses: list[dict[str, object]] = []
    seen_ids: set[str] = set()
    for row in rows[header_index + 1 :]:
        cells = list(row) + [None] * (6 - len(row))
        if all(value is None or _text(value) == "" for value in cells[:6]):
            break
        number, clause_id, kind, weight, share, description = cells[:6]
        clause_id, kind, description = _text(clause_id), _text(kind), _text(description)
        label = f"{name}: clause {number!r}"
        if not CLAUSE_ID_PATTERN.match(clause_id):
            problems.append(f"{label} ID {clause_id!r} must be lower-case snake_case")
        elif clause_id in seen_ids:
            problems.append(f"{label} ID {clause_id!r} is duplicated")
        seen_ids.add(clause_id)
        if kind not in KINDS:
            problems.append(f"{label} kind {kind!r} must be 'qualitative' or 'numeric'")
        if not _is_positive_number(weight):
            problems.append(f"{label} weight {weight!r} must be a positive number")
        elif float(weight).is_integer():
            weight = int(weight)
        if not description:
            problems.append(f"{label} has no criterion text")
        clauses.append({"id": clause_id, "kind": kind, "weight": weight, "listed_share": share, "description": description})
    if not clauses:
        problems.append(f"{name}: clause table is empty")
        return None
    return clauses


def build_registry(workbook_path: str | Path, registry: dict[str, object]) -> dict[str, object]:
    """Rebuild the registry from the workbook, keeping each paper's stem aliases from the current registry."""
    version, workbook_criteria = read_criteria_workbook(workbook_path)
    stems_by_paper = {}
    for criterion in registry.get("criteria", []):
        stems = criterion.get("paper_stems") if isinstance(criterion, dict) else None
        if not isinstance(stems, list) or not stems or not isinstance(stems[0], str):
            raise ValueError("Each registry criterion must list paper_stems with the canonical stem first.")
        stems_by_paper[stems[0]] = stems
    missing = [criterion["paper_stem"] for criterion in workbook_criteria if criterion["paper_stem"] not in stems_by_paper]
    unused = sorted(set(stems_by_paper) - {criterion["paper_stem"] for criterion in workbook_criteria})
    if missing or unused:
        raise ValueError(
            "Workbook sheets and registry papers differ"
            + (f"; add paper_stems for: {', '.join(missing)}" if missing else "")
            + (f"; no workbook sheet for: {', '.join(unused)}" if unused else "")
        )
    return {
        "version": version,
        "source": {"workbook": Path(workbook_path).name, "authoritative": "clause table of each paper sheet"},
        "criteria": [
            {
                "paper_stems": stems_by_paper[criterion["paper_stem"]],
                "target": criterion["target"],
                "scoring_instructions": criterion["scoring_instructions"],
                "clauses": criterion["clauses"],
            }
            for criterion in workbook_criteria
        ],
    }


def _is_number(value: object) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _is_positive_number(value: object) -> bool:
    return _is_number(value) and value > 0


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def formula_clauses(formula: str) -> list[str] | None:
    """Clause texts of a '1/n*( a + b )' formula with whitespace normalized; [] for an empty cell."""
    if not formula:
        return []
    match = FORMULA_PATTERN.match(formula.strip())
    if match is None:
        return None
    clauses = [_normalized(clause) for clause in match.group("body").split(" + ")]
    return clauses if len(clauses) == int(match.group("count")) else None


def _normalized(value: str) -> str:
    return " ".join(value.split())


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--registry", type=Path, default=DEFAULT_REGISTRY)
    parser.add_argument("--update", action="store_true", help="Write the synchronized registry.")
    args = parser.parse_args(argv)

    registry = json.loads(args.registry.read_text(encoding="utf8"))
    try:
        synchronized = build_registry(args.workbook, registry)
    except CriteriaWorkbookError as error:
        print(error)
        return 2
    if args.update:
        args.registry.write_text(json.dumps(synchronized, indent=2, ensure_ascii=False) + "\n", encoding="utf8")
        print(f"Synchronized {args.registry} with {args.workbook} (criteria v{synchronized['version']}).")
        return 0
    if synchronized != registry:
        print(f"{args.registry} is out of sync with {args.workbook}; rerun with --update.")
        return 1
    print(f"{args.registry} matches {args.workbook}.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
