import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

import manual_biological_scores as manual_scores


HEADERS = [
    "No|mode|replicate",
    "Paper",
    "Biological Domain",
    "Model Formalism",
    "Authors' Simulator",
    "Please focus on",
    "Result Fig.",
    "Fig. License",
    None,
    "Result Criterion",
    "Tech score by 11.05.2026",
    "Attempt Count",
    "Mode",
    "Result by 31.03.2026 (Prerana)",
    "14.04.2026 (Simon)",
    "20.04.2026 (Lutz)",
    "28.04.2026 (Lutz)",
    "Agent's selfeval",
    "11.05.2026 (Lutz) with broken _mcp script",
    "01.06.2026 (now full papers): Lutz",
    "01.06.2026 Prerana+Anjali",
    "01.06.2026 Mean Score",
    "Diff_scores",
    "Latest result comment by Lutz",
    "Latest result comment by Prerana and Anjali",
    None,
    None,
    "Criterion 1 (qualitative)",
    "Criterion 2 (quantitative)",
    "Criterion 1 Scores",
    "Criterion 2 Scores",
    "Mean",
]


class ManualBiologicalScoresTests(unittest.TestCase):
    def test_parse_workbook_run_id_normalizes_valid_ids(self) -> None:
        self.assertEqual(manual_scores.parse_workbook_run_id("11a1 -> 15a1"), (11, "a", 1, "11a1"))
        self.assertEqual(manual_scores.parse_workbook_run_id(" 2D03 "), (2, "d", 3, "2d3"))
        self.assertIsNone(manual_scores.parse_workbook_run_id(""))
        self.assertIsNone(manual_scores.parse_workbook_run_id("paper"))

    def test_workbook_and_string_helpers_handle_sparse_inputs(self) -> None:
        with TemporaryDirectory() as tmp:
            workbook_path = Path(tmp) / "manual.xlsx"
            write_workbook(
                workbook_path,
                [
                    ["1a1", "Szabó et al. (2010)", "python_web", None, None, None, 1.0],
                    ["2b1", "Additional biology papers", "morpheus_cli", None, None, None, 1.0],
                    ["not-a-run", "Paper", "morpheus_cli", None, None, None, 1.0],
                ],
            )

            rows = manual_scores.workbook_rows(workbook_path)
            titles = manual_scores.group_titles(rows)

        self.assertEqual(len(rows), 3)
        self.assertEqual(titles, {1: "Szabó et al. (2010)"})
        self.assertEqual(manual_scores.paper_key_from_title("Szabó et al. (2010)"), "szabo2010")
        self.assertEqual(manual_scores.paper_key_from_title("No year title!"), "noyeartitle")
        self.assertEqual(manual_scores.paper_number_from_canonical_stem("15_Reza2025"), 15.0)
        self.assertTrue(math_is_nan(manual_scores.paper_number_from_canonical_stem("Reza2025")))
        self.assertTrue(math_is_nan(manual_scores.nullable_number(None)))
        self.assertTrue(math_is_nan(manual_scores.nullable_number("bad")))
        self.assertTrue(math_is_nan(manual_scores.nullable_number(float("inf"))))
        self.assertEqual(manual_scores.nullable_number("0.25"), 0.25)
        self.assertEqual(manual_scores.string_value(" value "), "value")
        self.assertEqual(manual_scores.string_value(3), "")
        self.assertEqual(manual_scores.cell(("a",), 2), None)
        self.assertEqual(manual_scores.strip_accents("Müller Szabó"), "Muller Szabo")
        self.assertEqual(
            manual_scores.default_output_paths(Path("scores/manual.xlsx"), Path("out")),
            (Path("out/manual_biological_scores.csv"), Path("out/manual_biological_scores_by_paper_mode.csv")),
        )

    def test_empty_workbooks_and_empty_aggregates_return_stable_schemas(self) -> None:
        with TemporaryDirectory() as tmp:
            workbook_path = Path(tmp) / "manual.xlsx"
            write_workbook(workbook_path, [])

            rows = manual_scores.extract_manual_score_rows(workbook_path)
            aggregate = manual_scores.aggregate_manual_scores(rows)

        self.assertEqual(list(rows.columns), manual_scores.ROW_COLUMNS)
        self.assertTrue(rows.empty)
        self.assertEqual(list(aggregate.columns), manual_scores.AGGREGATE_COLUMNS)
        self.assertTrue(aggregate.empty)

        with TemporaryDirectory() as tmp:
            workbook_path = Path(tmp) / "manual.xlsx"
            write_workbook(
                workbook_path,
                [
                    ["not-a-run", "Paper", "morpheus_cli", None, None, None, 1.0],
                    ["1a1", None, "python_web", None, None, None, 1.0],
                ],
            )

            rows = manual_scores.extract_manual_score_rows(workbook_path)

        self.assertEqual(list(rows.columns), manual_scores.ROW_COLUMNS)
        self.assertTrue(rows.empty)

    def test_normalize_mode_handles_aliases_and_parenthetical_suffixes(self) -> None:
        self.assertEqual(manual_scores.normalize_mode("morpheus_cli+docs"), "morpheus_cli_docs")
        self.assertEqual(manual_scores.normalize_mode("morpheus_cli+docs_offline"), "morpheus_cli_docs")
        self.assertEqual(manual_scores.normalize_mode("morpheus_cli_docs_offline"), "morpheus_cli_docs")
        self.assertEqual(manual_scores.normalize_mode("Morpheus CLI (offline)"), "morpheus_cli")
        self.assertEqual(manual_scores.normalize_mode(None), "")

    def test_manual_score_policy_uses_only_workbook_mean_score(self) -> None:
        score, source, count, overlap = manual_scores.manual_score_from_policy(0.666667)
        self.assertAlmostEqual(score, 0.666667)
        self.assertEqual((source, count, overlap), ("workbook_mean_score", 1, False))

        score, source, count, overlap = manual_scores.manual_score_from_policy(0.75)
        self.assertAlmostEqual(score, 0.75)
        self.assertEqual((source, count, overlap), ("workbook_mean_score", 1, False))

        score, source, count, overlap = manual_scores.manual_score_from_policy(0.9)
        self.assertAlmostEqual(score, 0.9)
        self.assertEqual((source, count, overlap), ("workbook_mean_score", 1, False))

        score, source, count, overlap = manual_scores.manual_score_from_policy(float("nan"))
        self.assertTrue(math_is_nan(score))
        self.assertEqual((source, count, overlap), ("missing_workbook_mean_score", 0, False))

    def test_papers_2_to_9_use_workbook_mean_score_column(self) -> None:
        with TemporaryDirectory() as tmp:
            workbook_path = Path(tmp) / "manual.xlsx"
            write_workbook(
                workbook_path,
                [
                    ["2a1", None, "python_web", None, 0.2, None, 0.9],
                    ["2b1", None, "morpheus_cli", None, None, 0.66, 0.66],
                    ["2d1", "Thapa et al. (2024)", "morpheus_mcp", None, None, 1.0, 1.0],
                ],
            )

            rows = manual_scores.extract_manual_score_rows(workbook_path)

            python_web = rows.loc[
                (rows["paper_key"] == "thapa2024")
                & (rows["mode"] == "python_web")
                & (rows["repetition"] == 1)
            ].iloc[0]
            self.assertEqual(python_web["paper_title"], "Thapa et al. (2024)")
            self.assertEqual(python_web["canonical_paper_stem"], "02_Thapa2024")
            self.assertAlmostEqual(python_web["manual_score"], 0.9)
            self.assertEqual(python_web["manual_score_source"], "workbook_mean_score")
            self.assertEqual(python_web["manual_score_source_count"], 1)
            self.assertFalse(python_web["manual_score_overlap"])
            self.assertAlmostEqual(python_web["workbook_mean_score"], 0.9)

            cli = rows.loc[
                (rows["paper_key"] == "thapa2024")
                & (rows["mode"] == "morpheus_cli")
                & (rows["repetition"] == 1)
            ].iloc[0]
            self.assertAlmostEqual(cli["manual_score"], 0.66)
            self.assertEqual(cli["manual_score_source"], "workbook_mean_score")

    def test_different_af_score_does_not_override_workbook_mean(self) -> None:
        with TemporaryDirectory() as tmp:
            workbook_path = Path(tmp) / "manual.xlsx"
            write_workbook(
                workbook_path,
                [
                    ["1d1", "Szabó et al. (2010)", "morpheus_mcp", None, 0.66, 0.67, 0.665, 0.75],
                    ["10d1", "Berkhout et al. (2025)", "morpheus_mcp", None, 0.67, 0.66, 0.75, 0.75],
                ],
            )

            rows = manual_scores.extract_manual_score_rows(workbook_path)

        szabo = rows.loc[rows["paper_key"] == "szabo2010"].iloc[0]
        self.assertAlmostEqual(szabo["manual_score"], 0.665)
        self.assertEqual(szabo["manual_score_source"], "workbook_mean_score")

        berkhout = rows.loc[rows["paper_key"] == "berkhout2025"].iloc[0]
        self.assertAlmostEqual(berkhout["manual_score"], 0.75)
        self.assertEqual(berkhout["manual_score_source"], "workbook_mean_score")

    def test_papers_2_to_9_use_workbook_mean_even_when_reviewer_columns_overlap(self) -> None:
        with TemporaryDirectory() as tmp:
            workbook_path = Path(tmp) / "manual.xlsx"
            write_workbook(
                workbook_path,
                [
                    ["3d1", "Carpenter et al. (2024)", "morpheus_mcp", None, 0.66, 1.0, 0.83],
                    ["3d2", None, "morpheus_mcp", None, 0.5, None, 0.5],
                ],
            )

            rows = manual_scores.extract_manual_score_rows(workbook_path)
            aggregate = manual_scores.aggregate_manual_scores(rows)

            overlap = rows.loc[rows["workbook_run_id"] == "3d1"].iloc[0]
            self.assertFalse(overlap["manual_score_overlap"])
            self.assertEqual(overlap["manual_score_source"], "workbook_mean_score")
            self.assertEqual(overlap["manual_score_source_count"], 1)
            self.assertAlmostEqual(overlap["manual_score"], 0.83)

            summary = aggregate.iloc[0]
            self.assertAlmostEqual(summary["manual_score_mean"], 0.665)
            self.assertEqual(summary["manual_score_count"], 2)
            self.assertEqual(summary["manual_score_overlap_count"], 0)

    def test_v3_pure_biology_ids_map_to_canonical_repo_stems_and_use_column_v(self) -> None:
        with TemporaryDirectory() as tmp:
            workbook_path = Path(tmp) / "manual.xlsx"
            write_workbook(
                workbook_path,
                [
                    ["11a1", None, "python_web", None, 0.25, None, 0.25],
                    ["11b1", None, "morpheus_cli", None, 0.75, None, 0.8],
                    ["11d1", "Anlas et al. (2024)", "morpheus_mcp", None, 0.75, None, 0.75],
                ],
            )

            rows = manual_scores.extract_manual_score_rows(workbook_path)

            cli = rows.loc[(rows["mode"] == "morpheus_cli") & (rows["repetition"] == 1)].iloc[0]
            self.assertEqual(cli["paper_key"], "anlas2024")
            self.assertEqual(cli["canonical_paper_stem"], "11_Anlas2024")
            self.assertEqual(cli["workbook_run_id"], "11b1")
            self.assertAlmostEqual(cli["manual_score"], 0.8)
            self.assertEqual(cli["manual_score_source"], "workbook_mean_score")
            self.assertFalse(cli["manual_score_overlap"])

    def test_updated_manual_scores_sheet_overrides_legacy_column_v(self) -> None:
        with TemporaryDirectory() as tmp:
            workbook_path = Path(tmp) / "manual.xlsx"
            write_workbook(
                workbook_path,
                [
                    ["1b1", None, "morpheus_cli", None, None, None, 0.5],
                    ["1d1", "Szabó et al. (2010)", "morpheus_mcp", None, None, None, 0.5],
                ],
                updated_rows=[
                    ["1b1", None, "morpheus_cli", 0.75],
                    ["1d1", "Szabó et al. (2010)", "morpheus_mcp", 1.0],
                    ["1e1", None, "morpheus_cli+docs (Offline)", None],
                ],
            )

            rows = manual_scores.extract_manual_score_rows(workbook_path)

        by_id = rows.set_index("workbook_run_id")
        self.assertAlmostEqual(by_id.loc["1b1", "manual_score"], 0.75)
        self.assertAlmostEqual(by_id.loc["1d1", "manual_score"], 1.0)
        self.assertEqual(by_id.loc["1b1", "manual_score_source"], "updated_manual_mean_score")
        self.assertEqual(by_id.loc["1d1", "result_criterion"], "space criterion | speed criterion")
        self.assertTrue(math_is_nan(by_id.loc["1e1", "manual_score"]))
        self.assertEqual(by_id.loc["1e1", "manual_score_source"], "missing_updated_manual_mean_score")
        self.assertEqual(by_id.loc["1e1", "mode"], "morpheus_cli_docs")
        self.assertTrue(math_is_nan(by_id.loc["1b1", "lutz_20260601_score"]))

    def test_updated_manual_scores_sheet_rejects_shifted_columns(self) -> None:
        with TemporaryDirectory() as tmp:
            workbook_path = Path(tmp) / "manual.xlsx"
            write_workbook(workbook_path, [], updated_rows=[["1a1", "Szabó et al. (2010)", "python_web", 0.0]])
            workbook = load_workbook(workbook_path)
            workbook["Updated_manual_scores_05.10.202"].cell(row=1, column=15, value="Comments")
            workbook.save(workbook_path)

            with self.assertRaisesRegex(ValueError, "layout changed"):
                manual_scores.extract_manual_score_rows(workbook_path)

    def test_export_writes_row_and_aggregate_csvs(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook_path = root / "manual.xlsx"
            rows_path = root / "manual_rows.csv"
            aggregate_path = root / "manual_aggregate.csv"
            write_workbook(
                workbook_path,
                [
                    ["1a1", None, "python_web", None, 0.2, None, None],
                    ["1d1", "Szabó et al. (2010)", "morpheus_mcp", None, None, 1.0, None],
                ],
            )

            written = manual_scores.export_manual_score_csvs(workbook_path, rows_path, aggregate_path)

            self.assertEqual(written["rows"], rows_path)
            self.assertEqual(written["aggregate"], aggregate_path)
            self.assertTrue(rows_path.exists())
            self.assertTrue(aggregate_path.exists())
            self.assertIn("canonical_paper_stem,manual_paper_number", rows_path.read_text(encoding="utf8"))

    def test_main_resolves_default_and_explicit_output_paths(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook_path = root / "manual.xlsx"
            workbook_path.write_text("placeholder", encoding="utf8")
            explicit_dir = root / "out"
            calls: list[tuple[Path, Path, Path]] = []

            def fake_export(workbook: str | Path, rows: str | Path, aggregate: str | Path) -> dict[str, Path]:
                calls.append((Path(workbook), Path(rows), Path(aggregate)))
                return {"rows": Path(rows), "aggregate": Path(aggregate)}

            with patch.object(manual_scores, "export_manual_score_csvs", side_effect=fake_export):
                default_status = manual_scores.main([str(workbook_path)])
                explicit_status = manual_scores.main([str(workbook_path), "--output-dir", str(explicit_dir)])

            self.assertEqual(default_status, 0)
            self.assertEqual(explicit_status, 0)
            self.assertEqual(calls[0], (
                workbook_path,
                workbook_path.with_name("manual_biological_scores.csv"),
                workbook_path.with_name("manual_biological_scores_by_paper_mode.csv"),
            ))
            self.assertEqual(calls[1], (
                workbook_path,
                explicit_dir / "manual_biological_scores.csv",
                explicit_dir / "manual_biological_scores_by_paper_mode.csv",
            ))


UPDATED_HEADERS = [
    "No|mode|replicate",
    "Paper",
    "Biological Domain",
    "Model Formalism",
    "Authors' Simulator",
    "Please focus on",
    "Result Fig.",
    "Fig. License",
    None,
    "Mode",
    "Criterion 1 (qualitative)",
    "Criterion 2 (quantitative)",
    "Criterion 1 Scores",
    "Criterion 2 Scores",
    "Mean",
    "Comments ",
]


def write_workbook(
    path: Path, rows: list[list[object]], updated_rows: list[list[object]] | None = None
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Leaderboard"
    sheet.append(HEADERS)
    for values in rows:
        run_id, paper, mode, may_score, lutz_score, pa_score, mean_score, *optional = values
        af_score = optional[0] if optional else None
        row = [None] * len(HEADERS)
        row[0] = run_id
        row[1] = paper
        row[9] = may_score
        row[12] = mode
        row[19] = lutz_score
        row[20] = pa_score
        row[21] = mean_score
        row[31] = af_score
        sheet.append(row)
    if updated_rows is not None:
        updated = workbook.create_sheet("Updated_manual_scores_05.10.202")
        updated.append(UPDATED_HEADERS)
        for run_id, paper, mode, mean_score in updated_rows:
            row = [None] * len(UPDATED_HEADERS)
            row[0] = run_id
            row[1] = paper
            row[9] = mode
            if paper:
                row[10] = "space criterion"
                row[11] = "speed criterion"
            row[14] = mean_score
            updated.append(row)
    workbook.save(path)


def math_is_nan(value: float) -> bool:
    return value != value


if __name__ == "__main__":
    unittest.main()
