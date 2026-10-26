import contextlib
import io
import json
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook

import sync_biological_scoring_criteria as criteria_sync


ROOT = Path(__file__).resolve().parents[1]

# Weights deliberately follow no rule: the script must take them from the table as written.
CLAUSES = [
    ("cells_present_end", "qualitative", 3, "cells are present at end"),
    ("pattern_visible_end", "qualitative", 1, "a pattern is visible at end"),
    ("ratio_between_2_and_3", "numeric", 1.5, "final/initial cell number is between 2 and 3"),
]


def write_workbook(
    path: Path,
    sheets: dict[str, list[tuple[str, str, float, str]]],
    version: int = 7,
    **overrides,
) -> None:
    """Write a criteria workbook in the per-paper layout; overrides patch cells as {(sheet, cell): value}."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.create_sheet("Export Summary")["B3"] = "This document was exported from Numbers."
    for stem, clauses in sheets.items():
        sheet = workbook.create_sheet(stem)
        sheet["A1"] = f"{stem}: biological scoring criteria v{version}"
        sheet["A3"], sheet["B3"] = "Target", f"Target for {stem}."
        sheet["A4"], sheet["B4"] = "Scoring instructions", "Score both categories with equal total weight."
        sheet["A5"] = "Workbook formula (qualitative)"
        sheet["B5"] = criteria_sync.criterion_formula([text for _, kind, _, text in clauses if kind == "qualitative"])
        sheet["A6"] = "Workbook formula (quantitative)"
        sheet["B6"] = criteria_sync.criterion_formula([text for _, kind, _, text in clauses if kind == "numeric"])
        sheet.append([])
        for column, header in zip("ABCDEF", ("#", "Clause ID", "Kind", "Weight", "Share of paper score", "Criterion")):
            sheet[f"{column}8"] = header
        total_weight = sum(weight for _, _, weight, _ in clauses)
        for index, (clause_id, kind, weight, text) in enumerate(clauses, 1):
            row = 8 + index
            for column, value in zip("ABCDEF", (index, clause_id, kind, weight, weight / total_weight, text)):
                sheet[f"{column}{row}"] = value
    for (stem, cell), value in overrides.get("cells", {}).items():
        workbook[stem][cell] = value
    workbook.save(path)


class SyncBiologicalScoringCriteriaTests(unittest.TestCase):
    def test_formula_clauses_tolerate_spacing_and_reject_malformed_formulas(self) -> None:
        self.assertEqual(criteria_sync.formula_clauses("1/2*(a  b + c )"), ["a b", "c"])
        self.assertEqual(criteria_sync.formula_clauses("1/2* ( a + c)"), ["a", "c"])
        self.assertEqual(criteria_sync.formula_clauses(""), [])
        self.assertIsNone(criteria_sync.formula_clauses("1/3*( a + c )"))
        self.assertIsNone(criteria_sync.formula_clauses("a + c"))

    def test_build_registry_reads_clause_tables_and_keeps_stem_aliases(self) -> None:
        with TemporaryDirectory() as tmp:
            workbook_path = Path(tmp) / "criteria.xlsx"
            write_workbook(workbook_path, {"01_Paper2020": CLAUSES})
            registry = {"version": 5, "criteria": [{"paper_stems": ["01_Paper2020", "Paper2020"], "clauses": []}]}

            synchronized = criteria_sync.build_registry(workbook_path, registry)

        self.assertEqual(synchronized["version"], 7)
        self.assertEqual(synchronized["source"]["workbook"], "criteria.xlsx")
        criterion = synchronized["criteria"][0]
        self.assertEqual(criterion["paper_stems"], ["01_Paper2020", "Paper2020"])
        self.assertEqual(criterion["target"], "Target for 01_Paper2020.")
        self.assertEqual(
            criterion["clauses"],
            [
                {"id": "cells_present_end", "kind": "qualitative", "weight": 3, "description": "cells are present at end"},
                {"id": "pattern_visible_end", "kind": "qualitative", "weight": 1, "description": "a pattern is visible at end"},
                {"id": "ratio_between_2_and_3", "kind": "numeric", "weight": 1.5, "description": "final/initial cell number is between 2 and 3"},
            ],
        )

    def test_read_reports_every_drift_between_table_formulas_weights_and_shares(self) -> None:
        with TemporaryDirectory() as tmp:
            workbook_path = Path(tmp) / "criteria.xlsx"
            write_workbook(
                workbook_path,
                {"01_Paper2020": CLAUSES, "02_Other2021": CLAUSES},
                cells={
                    ("01_Paper2020", "B5"): "1/2*( cells are present at end + an edited pattern clause )",
                    ("01_Paper2020", "B10"): "Pattern Visible",
                    ("01_Paper2020", "C11"): "Numeric",
                    ("01_Paper2020", "E9"): 0.5,
                    ("01_Paper2020", "F10"): None,
                    ("02_Other2021", "A1"): "02_Other2021: biological scoring criteria v6",
                    ("02_Other2021", "D10"): 0,
                },
            )

            with self.assertRaises(criteria_sync.CriteriaWorkbookError) as raised:
                criteria_sync.read_criteria_workbook(workbook_path)

        problems = "\n".join(raised.exception.problems)
        self.assertIn("ID 'Pattern Visible' must be lower-case snake_case", problems)
        self.assertIn("kind 'Numeric' must be 'qualitative' or 'numeric'", problems)
        self.assertIn("clause cells_present_end share is 0.5, but its weight gives 0.545455", problems)
        self.assertIn("02_Other2021: clause 2 weight 0 must be a positive number", problems)
        self.assertNotIn("02_Other2021: clause cells_present_end share", problems)
        self.assertIn("clause 2 has no criterion text", problems)
        self.assertIn("'Workbook formula (qualitative)' does not restate", problems)
        self.assertIn("sheets disagree on the criteria version: [6, 7]", problems)

    def test_build_registry_requires_matching_papers(self) -> None:
        with TemporaryDirectory() as tmp:
            workbook_path = Path(tmp) / "criteria.xlsx"
            write_workbook(workbook_path, {"01_Paper2020": CLAUSES})
            registry = {"criteria": [{"paper_stems": ["02_Other2021"]}]}

            with self.assertRaisesRegex(ValueError, "add paper_stems for: 01_Paper2020; no workbook sheet for: 02_Other2021"):
                criteria_sync.build_registry(workbook_path, registry)

    def test_main_checks_updates_and_rejects_inconsistent_workbooks(self) -> None:
        with TemporaryDirectory() as tmp:
            workbook_path = Path(tmp) / "criteria.xlsx"
            registry_path = Path(tmp) / "registry.json"
            write_workbook(workbook_path, {"01_Paper2020": CLAUSES})
            registry_path.write_text(json.dumps({"version": 6, "criteria": [{"paper_stems": ["01_Paper2020"]}]}), "utf8")
            args = ["--workbook", str(workbook_path), "--registry", str(registry_path)]

            with contextlib.redirect_stdout(io.StringIO()):
                self.assertEqual(criteria_sync.main(args), 1)
                self.assertEqual(criteria_sync.main([*args, "--update"]), 0)
                self.assertEqual(criteria_sync.main(args), 0)
            self.assertEqual(json.loads(registry_path.read_text("utf8"))["version"], 7)

            write_workbook(workbook_path, {"01_Paper2020": CLAUSES}, cells={("01_Paper2020", "C9"): "Qualitative"})
            output = io.StringIO()
            with contextlib.redirect_stdout(output):
                self.assertEqual(criteria_sync.main(args), 2)
            self.assertIn("kind 'Qualitative'", output.getvalue())

    def test_checked_in_registry_matches_checked_in_workbook(self) -> None:
        registry = json.loads((ROOT / criteria_sync.DEFAULT_REGISTRY).read_text(encoding="utf8"))

        synchronized = criteria_sync.build_registry(ROOT / criteria_sync.DEFAULT_WORKBOOK, registry)

        self.assertEqual(synchronized, registry)
        self.assertEqual(registry["version"], 7)
        self.assertEqual(len(registry["criteria"]), 15)


if __name__ == "__main__":
    unittest.main()
